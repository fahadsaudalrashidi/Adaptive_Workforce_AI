"""
Phase 36 — Automated HR Letter Generator tests.

Unlike every other module this session, this generator intentionally
includes real employee names in its output — that is its entire purpose.
The disciplines under test here are different in kind from the PII-masking
tests elsewhere in this suite:

  1. Correct data population from real fixture rows (name, ID, job title,
     department, hire date, salary all appear verbatim in the generated
     .docx, sourced only from the uploaded row — nothing fabricated).
  2. Strict numeric-salary honest degradation for Salary Certificates: a
     range-text salary value must produce `available=False` with a reason,
     never an approximated figure written into an official document.
  3. The End-of-Service Settlement letter's gratuity figure exactly matches
     calling labor_law_calculators.analyze_gratuity() directly for the same
     employee/row — the cross-link must reuse the real computed number, not
     recompute or duplicate the Article 84/85 formula.
  4. Every generated document contains the required draft/review disclaimer
     and the platform-generated footer note.
  5. Bulk (mail-merge) generation produces one correctly-personalized
     document per employee, skipping (with a stated reason, not a fabricated
     document) any employee a given letter type isn't available for.

Run: py -m pytest tests/test_hr_letter_generator.py -v
"""

from __future__ import annotations

import zipfile
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from docx import Document

from backend.core.hr_letter_generator import (
    DRAFT_REVIEW_MARK_EN,
    LETTER_TYPES,
    CompanyDetails,
    GratuityLetterInput,
    _bidi_isolate,
    _translate_department_for_letter,
    safe_record_id_for_lookup,
    build_bulk_letters,
    build_employment_certificate,
    build_end_of_service_letter,
    build_salary_certificate,
    employee_data_result_to_dict,
    extract_employee_letter_records,
    unavailable_employee_data_result,
)
from backend.core.labor_law_calculators import analyze_gratuity

_ROOT = Path(__file__).resolve().parent.parent
_FIXTURE = _ROOT / "hr_letter_generator_test_data.csv"
_NO_NAME_FIXTURE = _ROOT / "test_employee_data.csv"

_COMPANY = CompanyDetails(
    company_name="Al-Rawda Industries LLC",
    company_address="King Fahd Road, Riyadh 12345, Saudi Arabia",
    signatory_name="Ahmed Al-Fahad",
    signatory_title="Chief Human Resources Officer",
)


def _load_fixture() -> pd.DataFrame:
    return pd.read_csv(_FIXTURE)


def _docx_text(document_bytes: bytes) -> str:
    doc = Document(BytesIO(document_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


# ── Employee data extraction ────────────────────────────────────────────────

class TestExtractEmployeeLetterRecords:
    def test_real_fixture_extracts_usable_employees(self):
        df = _load_fixture()
        result = extract_employee_letter_records(df)
        assert result.available is True
        assert result.name_column == "Name"
        assert result.employee_id_column == "EmployeeID"
        assert result.job_title_column == "JobTitle"
        assert result.department_column == "Department"
        assert result.hire_date_column == "HireDate"
        assert result.salary_column == "MonthlySalary"
        # 10 rows total, 1 excluded (HLEMP008 has a missing Name)
        assert result.n_total_rows == 10
        assert result.n_usable == 9
        assert result.n_excluded_missing_required == 1

    def test_missing_employee_excluded_from_usable_set(self):
        df = _load_fixture()
        result = extract_employee_letter_records(df)
        ids = [e.employee_id for e in result.employees]
        assert "HLEMP008" not in ids

    def test_real_row_values_populated_verbatim(self):
        df = _load_fixture()
        result = extract_employee_letter_records(df)
        sara = next(e for e in result.employees if e.employee_id == "HLEMP001")
        assert sara.name == "Sara Al-Otaibi"
        assert sara.job_title == "Senior Analyst"
        assert sara.department == "Finance"
        assert sara.salary_sar == 12000.0

    def test_range_text_salary_excluded_not_approximated(self):
        df = _load_fixture()
        result = extract_employee_letter_records(df)
        khalid = next(e for e in result.employees if e.employee_id == "HLEMP004")
        assert khalid.salary_sar is None
        assert khalid.salary_raw_display == "From 8000 to 10000 SAR"
        assert result.n_excluded_invalid_salary == 1

    def test_missing_salary_cell_is_none_not_zero(self):
        df = _load_fixture()
        result = extract_employee_letter_records(df)
        layla = next(e for e in result.employees if e.employee_id == "HLEMP007")
        assert layla.salary_sar is None

    def test_dataset_without_name_or_hire_date_columns_is_unavailable(self):
        df = pd.read_csv(_NO_NAME_FIXTURE)
        result = extract_employee_letter_records(df)
        assert result.available is False
        assert result.columns_detected["name"] is False
        assert result.columns_detected["hire_date"] is False
        assert result.employees == []

    def test_empty_dataframe_unavailable(self):
        result = extract_employee_letter_records(pd.DataFrame())
        assert result.available is False

    def test_to_dict_serializes_without_error(self):
        df = _load_fixture()
        result = extract_employee_letter_records(df)
        d = employee_data_result_to_dict(result)
        assert d["available"] is True
        assert len(d["employees"]) == 9
        assert "excluded_letter_types_note_en" in d
        assert "warning" in d["excluded_letter_types_note_en"].lower() or "disciplinary" in d["excluded_letter_types_note_en"].lower()

    def test_unavailable_helper_never_raises_and_has_all_fields(self):
        r = unavailable_employee_data_result("no data")
        assert r.available is False
        assert r.employees == []


# ── Salary Certificate ──────────────────────────────────────────────────────

class TestSalaryCertificate:
    def _sara(self):
        df = _load_fixture()
        result = extract_employee_letter_records(df)
        return next(e for e in result.employees if e.employee_id == "HLEMP001")

    def _khalid(self):
        df = _load_fixture()
        result = extract_employee_letter_records(df)
        return next(e for e in result.employees if e.employee_id == "HLEMP004")

    def test_generates_for_employee_with_clean_salary(self):
        result = build_salary_certificate(self._sara(), _COMPANY, "en")
        assert result.available is True
        assert result.document_bytes is not None
        text = _docx_text(result.document_bytes)
        assert "Sara Al-Otaibi" in text
        assert "HLEMP001" in text
        assert "Senior Analyst" in text
        assert "Finance" in text
        assert "SAR 12,000.00" in text
        assert _COMPANY.company_name in text
        assert _COMPANY.signatory_name in text

    def test_range_text_salary_refuses_not_approximates(self):
        result = build_salary_certificate(self._khalid(), _COMPANY, "en")
        assert result.available is False
        assert result.document_bytes is None
        assert result.reason is not None
        assert "clean numeric" in result.reason.lower() or "range-text" in result.reason.lower() or "read" in result.reason.lower()

    def test_document_contains_draft_marking_and_footer(self):
        result = build_salary_certificate(self._sara(), _COMPANY, "en")
        text = _docx_text(result.document_bytes)
        assert DRAFT_REVIEW_MARK_EN in text
        assert "automatically generated" in text
        assert "not an original signed HR record" in text

    def test_arabic_language_produces_arabic_body(self):
        result = build_salary_certificate(self._sara(), _COMPANY, "ar")
        assert result.available is True
        text = _docx_text(result.document_bytes)
        assert "Sara Al-Otaibi" in text
        assert "خطاب تعريف بالراتب" in text
        assert "مسودة" in text


# ── Language correctness (Phase 38 regression — see PHASE_26 report) ───────
#
# Live testing reported the language selector appearing to have no effect
# on the generated letter, plus Arabic company text splicing awkwardly into
# English sentences. Investigation (unit tests, TestClient route calls, and
# a browser-level network interception of the actual multipart POST) could
# NOT reproduce "selecting Arabic still yields an all-English letter" — the
# frontend correctly sends language=ar and the engine correctly branches
# into the Arabic template. What WAS reproduced and fixed: genuine Word
# bidi-rendering breakage whenever a value in one script is interpolated
# inline into a sentence written in the other script. These tests lock in
# both directions of that fix (department translation + bidi isolation) so
# it cannot regress silently.

_ENGLISH_TEMPLATE_PHRASES = [
    "This is to certify that",
    "is currently employed with",
    "Sincerely,",
    "To Whom It May Concern",
    "Date:",
]
_ARABIC_TEMPLATE_PHRASES = [
    "إلى من يهمه الأمر",
    "مع خالص التقدير",
    "الرقم الوظيفي",
    "التاريخ:",
]


def _paragraph_is_rtl(paragraph) -> bool:
    """True if the paragraph's XML carries the w:bidi property python-docx
    doesn't expose a high-level accessor for — the same RTL marker
    _set_rtl_paragraph() in hr_letter_generator.py writes."""
    return "<w:bidi/>" in paragraph._p.xml or "w:bidi" in paragraph._p.xml


class TestLanguageCorrectness:
    def _sara(self):
        df = _load_fixture()
        result = extract_employee_letter_records(df)
        return next(e for e in result.employees if e.employee_id == "HLEMP001")

    def _mohammed(self):
        df = _load_fixture()
        result = extract_employee_letter_records(df)
        return next(e for e in result.employees if e.employee_id == "HLEMP002")

    def test_bidi_isolate_wraps_latin_text_for_arabic_paragraph(self):
        assert _bidi_isolate("Senior Analyst", "ar") == "⁦Senior Analyst⁩"

    def test_bidi_isolate_wraps_arabic_text_for_english_paragraph(self):
        assert _bidi_isolate("شركة الأفق", "en") == "⁧شركة الأفق⁩"

    def test_bidi_isolate_no_op_when_script_matches_language(self):
        assert _bidi_isolate("Senior Analyst", "en") == "Senior Analyst"
        assert _bidi_isolate("محلل أول", "ar") == "محلل أول"

    def test_bidi_isolate_empty_string_safe(self):
        assert _bidi_isolate("", "ar") == ""

    def test_translate_department_uses_curated_table(self):
        assert _translate_department_for_letter("Finance", "ar") == "مالية"
        assert _translate_department_for_letter("Human Resources", "ar") == "الموارد البشرية"

    def test_translate_department_falls_back_when_unmapped(self):
        # "Spaceship Operations" is not in EN_AR_DEPARTMENT_SYNONYMS — must
        # fall back to the original value, never fabricate a translation.
        assert _translate_department_for_letter("Spaceship Operations", "ar") == "Spaceship Operations"

    def test_translate_department_no_op_for_english_output(self):
        assert _translate_department_for_letter("Finance", "en") == "Finance"

    @pytest.mark.parametrize("letter_type", ["salary_certificate", "employment_certificate", "end_of_service_settlement"])
    def test_arabic_letter_contains_zero_english_template_phrases(self, letter_type):
        df = _load_fixture()
        emp = self._mohammed()
        if letter_type == "salary_certificate":
            r = build_salary_certificate(emp, _COMPANY, "ar")
        elif letter_type == "employment_certificate":
            r = build_employment_certificate(emp, _COMPANY, "ar")
        else:
            gratuity_result = analyze_gratuity(df)
            cand = next(c for c in gratuity_result.candidates if c.record_id_safe == safe_record_id_for_lookup(emp.employee_id))
            r = build_end_of_service_letter(
                emp, _COMPANY, "ar",
                gratuity_amount_sar=cand.applicable_amount_sar,
                gratuity_scenario=cand.applicable_scenario,
                years_of_service=cand.years_of_service,
                monthly_wage_sar=cand.monthly_wage_sar,
                wage_column_name=gratuity_result.salary_column,
                gratuity_computed_at=gratuity_result.computed_at,
            )
        assert r.available is True
        text = _docx_text(r.document_bytes)
        for phrase in _ENGLISH_TEMPLATE_PHRASES:
            assert phrase not in text, f"English template phrase {phrase!r} leaked into Arabic letter"
        for phrase in _ARABIC_TEMPLATE_PHRASES:
            assert phrase in text, f"Expected Arabic template phrase {phrase!r} missing"

    def test_english_letter_contains_zero_arabic_template_phrases(self):
        r = build_salary_certificate(self._sara(), _COMPANY, "en")
        text = _docx_text(r.document_bytes)
        for phrase in _ARABIC_TEMPLATE_PHRASES:
            assert phrase not in text

    def test_department_translated_not_left_in_english(self):
        """Sara's department is Finance — a curated table entry — so the
        Arabic letter must show the Arabic label, not the English one."""
        r = build_salary_certificate(self._sara(), _COMPANY, "ar")
        text = _docx_text(r.document_bytes)
        assert "مالية" in text
        assert "Finance" not in text

    def test_job_title_kept_untranslated_but_isolated(self):
        """Job titles are open free text this platform cannot reliably
        translate — kept verbatim, but wrapped in isolate marks so Word
        renders the direction change cleanly instead of splicing badly."""
        r = build_salary_certificate(self._sara(), _COMPANY, "ar")
        text = _docx_text(r.document_bytes)
        assert "Senior Analyst" in text
        assert "⁦Senior Analyst⁩" in text

    def test_arabic_company_name_isolated_inside_english_letter(self):
        """The reproduced bug case: an Arabic company name inline in an
        English sentence must be isolate-wrapped, not spliced raw."""
        company_ar = CompanyDetails(
            company_name="شركة الأفق للحلول التقنية", company_address="Riyadh",
            signatory_name="Ahmed Al-Fahad", signatory_title="CHRO",
        )
        r = build_salary_certificate(self._sara(), company_ar, "en")
        text = _docx_text(r.document_bytes)
        assert "⁧شركة الأفق للحلول التقنية⁩" in text

    def test_same_employee_language_toggle_produces_different_documents(self):
        emp = self._sara()
        r_en = build_salary_certificate(emp, _COMPANY, "en")
        r_ar = build_salary_certificate(emp, _COMPANY, "ar")
        text_en = _docx_text(r_en.document_bytes)
        text_ar = _docx_text(r_ar.document_bytes)
        assert text_en != text_ar
        assert "SALARY CERTIFICATE" in text_en and "SALARY CERTIFICATE" not in text_ar
        assert "خطاب تعريف بالراتب" in text_ar and "خطاب تعريف بالراتب" not in text_en

    def test_arabic_body_paragraphs_render_rtl(self):
        r = build_salary_certificate(self._sara(), _COMPANY, "ar")
        doc = Document(BytesIO(r.document_bytes))
        body_paragraphs = [p for p in doc.paragraphs if "نشهد بأن الموظف" in p.text]
        assert body_paragraphs, "Arabic body paragraph not found"
        assert _paragraph_is_rtl(body_paragraphs[0])

    def test_english_body_paragraphs_not_marked_rtl(self):
        r = build_salary_certificate(self._sara(), _COMPANY, "en")
        doc = Document(BytesIO(r.document_bytes))
        body_paragraphs = [p for p in doc.paragraphs if "This is to certify that" in p.text]
        assert body_paragraphs
        assert not _paragraph_is_rtl(body_paragraphs[0])


# ── Employment Certificate ──────────────────────────────────────────────────

class TestEmploymentCertificate:
    def _mohammed(self):
        df = _load_fixture()
        result = extract_employee_letter_records(df)
        return next(e for e in result.employees if e.employee_id == "HLEMP002")

    def test_generates_without_needing_salary(self):
        result = build_employment_certificate(self._mohammed(), _COMPANY, "en")
        assert result.available is True
        text = _docx_text(result.document_bytes)
        assert "Mohammed Al-Qahtani" in text
        assert "HR Manager" in text
        assert "Human Resources" in text
        assert DRAFT_REVIEW_MARK_EN in text

    def test_no_salary_figure_leaked_into_employment_certificate(self):
        result = build_employment_certificate(self._mohammed(), _COMPANY, "en")
        text = _docx_text(result.document_bytes)
        assert "18,000" not in text


# ── End-of-Service Settlement — cross-link verification ────────────────────

class TestEndOfServiceCrossLink:
    def test_gratuity_figure_matches_labor_law_calculator_exactly(self):
        """The central cross-link claim: the letter's gratuity figure must
        be bit-for-bit the same number analyze_gratuity() computes directly
        for the same employee — never recomputed or approximated here."""
        df = _load_fixture()
        gratuity_result = analyze_gratuity(df)
        assert gratuity_result.available is True

        letter_result = extract_employee_letter_records(df)
        reem = next(e for e in letter_result.employees if e.employee_id == "HLEMP010")

        target_hash = safe_record_id_for_lookup(reem.employee_id)
        candidate = next(c for c in gratuity_result.candidates if c.record_id_safe == target_hash)

        # Reem: 12 years, resignation, 10y+ tier => full award (Article 85)
        assert candidate.applicable_scenario == "resignation"
        assert candidate.applicable_amount_sar == candidate.gratuity_full_award_sar

        built = build_end_of_service_letter(
            reem, _COMPANY, "en",
            gratuity_amount_sar=candidate.applicable_amount_sar,
            gratuity_scenario=candidate.applicable_scenario,
            years_of_service=candidate.years_of_service,
            monthly_wage_sar=candidate.monthly_wage_sar,
            wage_column_name=gratuity_result.salary_column,
            gratuity_computed_at=gratuity_result.computed_at,
        )
        assert built.available is True
        text = _docx_text(built.document_bytes)
        expected_amount_str = f"SAR {candidate.applicable_amount_sar:,.2f}"
        assert expected_amount_str in text
        assert "Reem Al-Subaie" in text
        assert "Article" in text

    def test_letter_never_computes_gratuity_itself(self):
        """build_end_of_service_letter must not import or call analyze_gratuity —
        confirmed by checking the module does not reference labor_law_calculators
        at all (route-layer decoupling, see module docstring)."""
        import backend.core.hr_letter_generator as mod
        assert "labor_law_calculators" not in mod.__dict__
        import inspect
        src = inspect.getsource(mod)
        assert "import backend.core.labor_law_calculators" not in src
        assert "from backend.core.labor_law_calculators" not in src

    def test_declining_scenario_when_gratuity_not_computable(self):
        """HLEMP004 (Khalid) has a range-text salary — gratuity cannot be
        computed for him, so the route layer must never invent a figure."""
        df = _load_fixture()
        gratuity_result = analyze_gratuity(df)
        khalid_hash = safe_record_id_for_lookup("HLEMP004")
        assert not any(c.record_id_safe == khalid_hash for c in gratuity_result.candidates)


# ── Draft marking on every letter type ──────────────────────────────────────

class TestDraftMarkingUniversal:
    @pytest.mark.parametrize("letter_type", LETTER_TYPES)
    def test_every_letter_type_carries_draft_marking(self, letter_type):
        df = _load_fixture()
        employees = extract_employee_letter_records(df).employees
        emp = next(e for e in employees if e.employee_id == "HLEMP002")  # clean salary + years

        if letter_type == "salary_certificate":
            r = build_salary_certificate(emp, _COMPANY, "en")
        elif letter_type == "employment_certificate":
            r = build_employment_certificate(emp, _COMPANY, "en")
        else:
            gratuity_result = analyze_gratuity(df)
            cand = next(c for c in gratuity_result.candidates if c.record_id_safe == safe_record_id_for_lookup(emp.employee_id))
            r = build_end_of_service_letter(
                emp, _COMPANY, "en",
                gratuity_amount_sar=cand.applicable_amount_sar,
                gratuity_scenario=cand.applicable_scenario,
                years_of_service=cand.years_of_service,
                monthly_wage_sar=cand.monthly_wage_sar,
                wage_column_name=gratuity_result.salary_column,
                gratuity_computed_at=gratuity_result.computed_at,
            )
        assert r.available is True
        text = _docx_text(r.document_bytes)
        assert DRAFT_REVIEW_MARK_EN in text


# ── Bulk (mail-merge) generation ────────────────────────────────────────────

class TestBulkGeneration:
    def test_bulk_salary_certificates_personalized_and_skip_bad_salary(self):
        df = _load_fixture()
        employees = extract_employee_letter_records(df).employees
        result = build_bulk_letters(employees, "salary_certificate", _COMPANY, "en")

        assert result.n_requested == len(employees)
        # Khalid (range-text salary) and Layla (missing salary) must be skipped
        khalid_item = next(i for i in result.items if i.employee_id == "HLEMP004")
        assert khalid_item.generated is False
        assert khalid_item.reason is not None

        assert result.n_generated > 0
        assert result.zip_bytes is not None

        with zipfile.ZipFile(BytesIO(result.zip_bytes)) as zf:
            names = zf.namelist()
            assert len(names) == result.n_generated
            # Verify at least two different documents are genuinely different content
            sara_name = next(n for n in names if "HLEMP001" in n)
            mohammed_name = next(n for n in names if "HLEMP002" in n)
            sara_text = _docx_text(zf.read(sara_name))
            mohammed_text = _docx_text(zf.read(mohammed_name))
            assert "Sara Al-Otaibi" in sara_text
            assert "Mohammed Al-Qahtani" not in sara_text
            assert "Mohammed Al-Qahtani" in mohammed_text
            assert sara_text != mohammed_text

    def test_bulk_employment_certificates_all_generated(self):
        df = _load_fixture()
        employees = extract_employee_letter_records(df).employees
        result = build_bulk_letters(employees, "employment_certificate", _COMPANY, "en")
        # Employment certificates need no salary, so every extracted employee gets one
        assert result.n_generated == len(employees)
        assert result.n_skipped == 0

    def test_bulk_end_of_service_uses_gratuity_lookup_and_skips_unavailable(self):
        df = _load_fixture()
        employees = extract_employee_letter_records(df).employees
        gratuity_result = analyze_gratuity(df)
        lookup = {}
        for cand in gratuity_result.candidates:
            for emp in employees:
                if safe_record_id_for_lookup(emp.employee_id) == cand.record_id_safe:
                    lookup[emp.employee_id] = GratuityLetterInput(
                        gratuity_amount_sar=cand.applicable_amount_sar,
                        gratuity_scenario=cand.applicable_scenario,
                        years_of_service=cand.years_of_service,
                        monthly_wage_sar=cand.monthly_wage_sar,
                        wage_column_name=gratuity_result.salary_column,
                        gratuity_computed_at=gratuity_result.computed_at,
                    )

        result = build_bulk_letters(employees, "end_of_service_settlement", _COMPANY, "en", gratuity_lookup=lookup)
        khalid_item = next(i for i in result.items if i.employee_id == "HLEMP004")
        assert khalid_item.generated is False  # range-text salary => no gratuity possible
        assert result.n_generated == len(lookup)

    def test_bulk_zero_generated_returns_none_zip(self):
        df = _load_fixture()
        employees = extract_employee_letter_records(df).employees
        result = build_bulk_letters(employees, "end_of_service_settlement", _COMPANY, "en", gratuity_lookup={})
        assert result.n_generated == 0
        assert result.zip_bytes is None
