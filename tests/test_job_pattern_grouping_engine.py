"""
Phase 19 — Job Pattern Grouping & Market Research Export engine tests.
Phase 20 — 5th grouping dimension: academic qualification.

Covers, each with real computed-value assertions (not smoke-only tests):
  1. Job-title normalization: mixed Arabic/English/casing variants of the
     same title merge correctly ("Manager" / "manager" / "مدير"), while a
     genuine near-duplicate typo ("Sales Exective" vs "Sales Executive") is
     flagged for human review but NEVER auto-merged — same discipline as
     department_normalizer.py.
  2. Hand-computed average/median/min/max salary per pattern.
  3. Honest partial-capability degradation: a dataset with job title +
     department but NO experience column still groups (by fewer
     dimensions), disclosing exactly which dimension is missing rather than
     silently dropping it. Phase 20 adds the same check for a dataset with
     no academic-qualification column at all — the 5th dimension must be
     just as optional as the other three.
  4. "Not detected" when no job-title column exists at all — the expected
     result for most fixtures already in this repo, none of which carry a
     dedicated job-title column (test_employee_data.csv's JobRole is the
     one exception, used here as the real-fixture positive case instead).
  5. Salary suppression: patterns with fewer than
     MIN_PATTERN_SIZE_FOR_SALARY_DISCLOSURE employees never show a salary
     figure — a mean of 1-2 people is not an aggregate, it is those
     people's literal pay.
  6. The Excel export contains no employee names or IDs anywhere in either
     sheet — aggregate pattern-level data only.
  7. Phase 20: academic-qualification normalization (degree-level merging
     across whitespace/EN/AR variants), the composite "degree — field"
     label when both a degree and a separate field/major column are
     present, finer-grained splitting vs. the pre-Phase-20 4-dimension
     grouping on the exact same rows, and a regression test for a real bug
     found while building this phase (a genuinely missing job title was
     silently fabricated into its own "Unknown / Unspecified Job Title"
     pattern instead of being excluded — see
     job_pattern_grouping_engine.py's _MISSING_TITLE_KEY comment).

Run: py -m pytest tests/test_job_pattern_grouping_engine.py -v
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from backend.core.job_pattern_grouping_engine import (
    MIN_PATTERN_SIZE_FOR_SALARY_DISCLOSURE,
    analyze_job_pattern_grouping,
    generate_market_research_workbook,
    job_pattern_grouping_from_dict,
    job_pattern_grouping_to_dict,
    normalize_degree_series,
    normalize_title_series,
    unavailable_job_pattern_result,
)

_ROOT = Path(__file__).resolve().parent.parent
_TITLE_FIXTURE = _ROOT / "job_pattern_title_test_data.csv"
_NO_EXPERIENCE_FIXTURE = _ROOT / "job_pattern_no_experience_test_data.csv"
_QUALIFICATION_FIXTURE = _ROOT / "job_pattern_qualification_test_data.csv"
_REAL_FIXTURE = _ROOT / "test_employee_data.csv"


# ── 1. Job-title normalization: merge known variants, flag near-dupes ──────

def test_title_variants_merge_across_casing_and_script():
    df = pd.read_csv(_TITLE_FIXTURE, encoding="utf-8")
    result = normalize_title_series(df["JobTitle"].astype(object))
    keys = result.normalized_keys
    # Manager / manager / مدير must all normalize to the same key.
    manager_keys = {keys.iloc[0], keys.iloc[1], keys.iloc[2]}
    assert len(manager_keys) == 1
    assert result.display_labels[keys.iloc[0]] == "Manager"


def test_near_duplicate_title_flagged_not_merged():
    df = pd.read_csv(_TITLE_FIXTURE, encoding="utf-8")
    result = normalize_title_series(df["JobTitle"].astype(object))
    # "Sales Executive" (row 3) and "Sales Exective" (row 4, typo) must be
    # DIFFERENT normalized keys...
    assert df["JobTitle"].iloc[3] == "Sales Executive"
    assert df["JobTitle"].iloc[4] == "Sales Exective"
    keys = result.normalized_keys
    assert keys.iloc[3] != keys.iloc[4]
    # ...but flagged together as a near-duplicate for human review.
    flat_flagged = {v for group in result.unmerged_similar_variants for v in group}
    assert "Sales Executive" in flat_flagged
    assert "Sales Exective" in flat_flagged


def test_engine_reports_title_merge_and_review_notes():
    df = pd.read_csv(_TITLE_FIXTURE, encoding="utf-8")
    result = analyze_job_pattern_grouping(df)
    assert result.available is True
    merge_notes = " ".join(result.title_normalization_notes)
    assert "Manager" in merge_notes
    assert "manager" in merge_notes
    review_notes = " ".join(result.title_normalization_notes)
    assert "Sales Exective" in review_notes and "Sales Executive" in review_notes


# ── 2. Hand-computed salary per pattern ─────────────────────────────────────

def test_hand_computed_salary_for_merged_manager_pattern():
    """E1/E2/E3 (Manager/manager/مدير, Sales, 7/8/9 years -> all Senior
    6-10yrs) have salaries 10000/12000/14000 -> avg=12000, median=12000,
    min=10000, max=14000, n=3 clears the disclosure floor."""
    df = pd.read_csv(_TITLE_FIXTURE, encoding="utf-8")
    result = analyze_job_pattern_grouping(df)
    manager_patterns = [p for p in result.patterns if p.job_title == "Manager"]
    assert len(manager_patterns) == 1
    p = manager_patterns[0]
    assert p.employee_count == 3
    assert p.n_with_valid_salary == 3
    assert p.salary_suppressed is False
    assert p.avg_salary_sar == 12000.0
    assert p.median_salary_sar == 12000.0
    assert p.salary_min_sar == 10000.0
    assert p.salary_max_sar == 14000.0
    assert p.experience_bucket == "Senior (6-10 yrs)"
    assert p.experience_range_label == "7-9 years"


def test_missing_experience_value_excluded_not_guessed():
    """E6 (Engineer, Engineering, YearsAtCompany missing, salary 9000) must
    form its own pattern with experience_bucket=None rather than being
    dropped or assigned a guessed bucket — distinct from E7 (Engineer,
    Engineering, 4 years -> Mid-Level)."""
    df = pd.read_csv(_TITLE_FIXTURE, encoding="utf-8")
    result = analyze_job_pattern_grouping(df)
    engineer_patterns = [p for p in result.patterns if p.job_title == "Engineer"]
    assert len(engineer_patterns) == 2
    buckets = {p.experience_bucket for p in engineer_patterns}
    assert None in buckets
    assert "Mid-Level (3-5 yrs)" in buckets


# ── 3. Honest partial-capability: no experience column at all ──────────────

def test_no_experience_column_still_groups_by_available_dimensions():
    df = pd.read_csv(_NO_EXPERIENCE_FIXTURE)
    result = analyze_job_pattern_grouping(df)
    assert result.available is True
    assert result.columns_detected["experience"] is False
    assert "experience" not in result.grouping_dimensions_used
    assert "job_title" in result.grouping_dimensions_used
    assert "department" in result.grouping_dimensions_used
    assert all(p.experience_bucket is None for p in result.patterns)
    assert "experience" in result.narrative_en.lower()


def test_no_experience_column_salary_still_computed_correctly():
    """3 Managers (10000/11000/12000) with no experience column present ->
    still groups into one Sales/Manager pattern with correct salary stats,
    proving the missing dimension degrades grouping granularity only, not
    salary computation."""
    df = pd.read_csv(_NO_EXPERIENCE_FIXTURE)
    result = analyze_job_pattern_grouping(df)
    manager_patterns = [p for p in result.patterns if p.job_title == "Manager"]
    assert len(manager_patterns) == 1
    p = manager_patterns[0]
    assert p.employee_count == 3
    assert p.avg_salary_sar == 11000.0
    assert p.salary_min_sar == 10000.0
    assert p.salary_max_sar == 12000.0


# ── 4. Not detected — no job-title column ───────────────────────────────────

def test_succession_fixture_has_no_job_title_column():
    df = pd.read_csv(_ROOT / "succession_test_data.csv")
    result = analyze_job_pattern_grouping(df)
    assert result.available is False
    assert result.columns_detected["job_title"] is False
    assert "job title" in result.reason.lower() or "job-title" in result.reason.lower()


def test_nitaqat_fixture_has_no_job_title_column():
    df = pd.read_csv(_ROOT / "nitaqat_test_data.csv")
    result = analyze_job_pattern_grouping(df)
    assert result.available is False


def test_unavailable_result_never_fabricates():
    result = unavailable_job_pattern_result(
        "No job-title column detected.",
        {"job_title": False, "job_level": False, "department": True, "experience": False, "salary": True},
    )
    assert result.available is False
    assert result.patterns == []
    assert result.n_patterns == 0
    d = job_pattern_grouping_to_dict(result)
    assert d["columns_detected"]["department"] is True
    assert d["columns_detected"]["job_title"] is False


# ── 5. Real fixture positive case (test_employee_data.csv's JobRole) ───────

def test_real_fixture_with_job_role_column_produces_patterns():
    df = pd.read_csv(_REAL_FIXTURE)
    result = analyze_job_pattern_grouping(df)
    assert result.available is True
    assert result.job_title_column == "JobRole"
    assert result.department_column == "Department"
    assert result.experience_column == "YearsAtCompany"
    assert result.salary_column == "MonthlyIncome"
    assert result.n_employees_total == 200
    assert result.n_employees_grouped == 200
    assert result.n_patterns > 0
    assert result.n_patterns < 200  # genuine grouping happened, not 1 pattern per row
    # Every JobRole in the source data is a known value with no missing rows.
    for p in result.patterns:
        assert p.employee_count > 0


def test_round_trip_serialization_preserves_patterns():
    df = pd.read_csv(_REAL_FIXTURE)
    result = analyze_job_pattern_grouping(df)
    d = job_pattern_grouping_to_dict(result)
    restored = job_pattern_grouping_from_dict(d)
    assert restored.available == result.available
    assert restored.n_patterns == result.n_patterns
    assert len(restored.patterns) == len(result.patterns)
    assert restored.patterns[0].avg_salary_sar == result.patterns[0].avg_salary_sar


# ── 6. Excel export: aggregate pattern-level data only, no PII ─────────────

def test_excel_export_contains_no_employee_ids_or_names():
    df = pd.read_csv(_TITLE_FIXTURE, encoding="utf-8")
    result = analyze_job_pattern_grouping(df)
    xlsx_bytes = generate_market_research_workbook(result)

    import io
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert set(wb.sheetnames) == {"Instructions", "Job Patterns"}

    raw_employee_ids = set(df["EmployeeID"].astype(str))
    all_cell_values = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    all_cell_values.append(str(cell.value))
    combined_text = " | ".join(all_cell_values)

    for emp_id in raw_employee_ids:
        assert emp_id not in combined_text, f"Employee ID {emp_id!r} leaked into export"


def test_excel_export_has_expected_columns():
    df = pd.read_csv(_REAL_FIXTURE)
    result = analyze_job_pattern_grouping(df)
    xlsx_bytes = generate_market_research_workbook(result)

    import io
    patterns_df = pd.read_excel(io.BytesIO(xlsx_bytes), sheet_name="Job Patterns")
    expected_cols = [
        "Pattern ID", "Job Title (normalized)", "Level", "Department",
        "Academic Qualification", "Experience Range", "Employee Count",
        "Current Avg Salary (SAR)", "Current Salary Range (SAR)",
    ]
    assert list(patterns_df.columns) == expected_cols
    assert len(patterns_df) == result.n_patterns


def test_excel_export_suppresses_small_pattern_salary():
    df = pd.read_csv(_TITLE_FIXTURE, encoding="utf-8")
    result = analyze_job_pattern_grouping(df)
    xlsx_bytes = generate_market_research_workbook(result)

    import io
    patterns_df = pd.read_excel(io.BytesIO(xlsx_bytes), sheet_name="Job Patterns")
    small_pattern_rows = patterns_df[patterns_df["Employee Count"] < MIN_PATTERN_SIZE_FOR_SALARY_DISCLOSURE]
    assert len(small_pattern_rows) > 0
    assert all(small_pattern_rows["Current Salary Range (SAR)"] == "Suppressed (n<3)")


# ── 7. Phase 20 — academic qualification (5th dimension) ───────────────────

def test_degree_variants_merge_across_whitespace_and_script():
    """'Bachelor's' / ' Bachelor's' (leading space, the exact real-world
    variant found in the Original Dataset of Employee Attrition) and
    'Ph.D' / 'PhD' / 'دكتوراه' must each normalize to one canonical key."""
    df = pd.read_csv(_QUALIFICATION_FIXTURE, encoding="utf-8")
    result = normalize_degree_series(df["Academic_degree"].astype(object))
    keys = result.normalized_keys
    bachelor_rows = df.index[df["JobTitle"] == "Sales Executive"].tolist()[:3]
    bachelor_keys = {keys.loc[i] for i in bachelor_rows}
    assert len(bachelor_keys) == 1
    assert result.display_labels[keys.loc[bachelor_rows[0]]] == "Bachelor's"

    phd_rows = df.index[df["JobTitle"] == "Engineer"].tolist()
    phd_keys = {keys.loc[i] for i in phd_rows}
    assert len(phd_keys) == 1
    assert result.display_labels[keys.loc[phd_rows[0]]] == "PhD"


def test_qualification_column_detected_and_composite_label_built():
    """When both Academic_degree and a separate EducationField column are
    present, the pattern's academic_qualification is the composite
    'Degree — Field' label, not just the degree alone."""
    df = pd.read_csv(_QUALIFICATION_FIXTURE, encoding="utf-8")
    result = analyze_job_pattern_grouping(df)
    assert result.available is True
    assert result.academic_degree_column == "Academic_degree"
    assert result.education_field_column == "EducationField"
    assert "academic_qualification" in result.grouping_dimensions_used

    marketing_pattern = next(p for p in result.patterns if p.academic_qualification == "Bachelor's — Marketing")
    assert marketing_pattern.employee_count == 3
    assert marketing_pattern.avg_salary_sar == 8500.0
    assert marketing_pattern.salary_suppressed is False


def test_qualification_dimension_produces_finer_grained_split():
    """The brief's own motivating example: two Sales Executives in the same
    department with the same experience band but DIFFERENT degree fields
    (Marketing vs. Engineering) must land in different patterns once the
    qualification dimension is active — proving the 5th dimension actually
    changes grouping outcomes, not just adds an unused column."""
    df = pd.read_csv(_QUALIFICATION_FIXTURE, encoding="utf-8")
    result = analyze_job_pattern_grouping(df)
    sales_patterns = [p for p in result.patterns if p.job_title == "Sales Executive"]
    assert len(sales_patterns) == 2
    qualifications = {p.academic_qualification for p in sales_patterns}
    assert qualifications == {"Bachelor's — Marketing", "Bachelor's — Engineering"}
    # The smaller (Engineering, n=2) sub-pattern must correctly fall under
    # the salary-disclosure floor while the larger (Marketing, n=3) one
    # does not — proving the floor interacts correctly with the finer split.
    marketing = next(p for p in sales_patterns if p.academic_qualification == "Bachelor's — Marketing")
    engineering = next(p for p in sales_patterns if p.academic_qualification == "Bachelor's — Engineering")
    assert marketing.employee_count == 3 and marketing.salary_suppressed is False
    assert engineering.employee_count == 2 and engineering.salary_suppressed is True


def test_missing_qualification_value_excluded_not_guessed():
    """E9 (Analyst, Operations, both Academic_degree and EducationField
    blank) must have academic_qualification=None rather than being dropped
    or assigned a fabricated value."""
    df = pd.read_csv(_QUALIFICATION_FIXTURE, encoding="utf-8")
    result = analyze_job_pattern_grouping(df)
    analyst_patterns = [p for p in result.patterns if p.job_title == "Analyst"]
    assert len(analyst_patterns) == 1
    assert analyst_patterns[0].academic_qualification is None
    assert analyst_patterns[0].employee_count == 1


def test_no_qualification_column_falls_back_to_four_dimensions():
    """test_employee_data.csv has no Academic_degree column at all — the
    5th dimension must be simply omitted (same graceful degradation as the
    other 3 optional dimensions), not block the whole feature."""
    df = pd.read_csv(_REAL_FIXTURE)
    result = analyze_job_pattern_grouping(df)
    assert result.available is True
    assert result.columns_detected["academic_degree"] is False
    assert result.columns_detected["education_field"] is False
    assert "academic_qualification" not in result.grouping_dimensions_used
    assert all(p.academic_qualification is None for p in result.patterns)
    assert "academic_qualification" in result.narrative_en.lower()


def test_five_dimension_grouping_is_finer_than_four_dimension_on_same_rows():
    """Building the qualification fixture's patterns WITHOUT the
    qualification dimension (by dropping the two qualification columns
    before analysis) must produce STRICTLY FEWER patterns than analyzing
    the same rows WITH qualification — direct proof the 5th dimension
    subdivides existing groups rather than just adding an inert column."""
    df_with = pd.read_csv(_QUALIFICATION_FIXTURE, encoding="utf-8")
    df_without = df_with.drop(columns=["Academic_degree", "EducationField"])

    result_with = analyze_job_pattern_grouping(df_with)
    result_without = analyze_job_pattern_grouping(df_without)

    assert result_without.available is True
    assert "academic_qualification" not in result_without.grouping_dimensions_used
    assert result_with.n_patterns > result_without.n_patterns
    # Specifically: the 2 Sales Executive patterns collapse into 1 without
    # qualification data to split them.
    sales_without = [p for p in result_without.patterns if p.job_title == "Sales Executive"]
    assert len(sales_without) == 1
    assert sales_without[0].employee_count == 5


def test_missing_job_title_row_excluded_not_fabricated_as_pattern():
    """Regression test for a real bug found while building this phase: the
    missing-value sentinel (MISSING_TITLE_LABEL) passes through the same
    exact-normalize (trim+casefold) tier as every real title before
    becoming a grouping key, so comparing a row's key against the
    ORIGINAL-CASE constant silently failed to exclude it — a row with a
    genuinely missing job title was fabricated into its own 'Unknown /
    Unspecified Job Title' pattern instead of being excluded. Fixed by
    comparing against the casefolded _MISSING_TITLE_KEY instead."""
    df = pd.DataFrame({
        "EmployeeID": ["E1", "E2"],
        "JobTitle": ["Manager", None],
        "Department": ["Sales", "Sales"],
    })
    result = analyze_job_pattern_grouping(df)
    assert result.n_employees_grouped == 1
    assert result.n_patterns == 1
    assert result.patterns[0].job_title == "Manager"
    assert result.patterns[0].employee_count == 1


def test_excel_export_includes_academic_qualification_column_and_values():
    df = pd.read_csv(_QUALIFICATION_FIXTURE, encoding="utf-8")
    result = analyze_job_pattern_grouping(df)
    xlsx_bytes = generate_market_research_workbook(result)

    import io
    patterns_df = pd.read_excel(io.BytesIO(xlsx_bytes), sheet_name="Job Patterns")
    assert "Academic Qualification" in patterns_df.columns
    values = set(patterns_df["Academic Qualification"].fillna(""))
    assert "Bachelor's — Marketing" in values
    assert "Bachelor's — Engineering" in values
    assert "PhD" in values


def test_excel_export_still_contains_no_employee_ids_with_qualification_column():
    """Re-verification of Phase 19's PII-safety guarantee now that a 5th
    column has been added — must still hold."""
    df = pd.read_csv(_QUALIFICATION_FIXTURE, encoding="utf-8")
    result = analyze_job_pattern_grouping(df)
    xlsx_bytes = generate_market_research_workbook(result)

    import io
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    raw_employee_ids = set(df["EmployeeID"].astype(str))
    all_cell_values = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    all_cell_values.append(str(cell.value))
    combined_text = " | ".join(all_cell_values)
    for emp_id in raw_employee_ids:
        assert emp_id not in combined_text, f"Employee ID {emp_id!r} leaked into export"
