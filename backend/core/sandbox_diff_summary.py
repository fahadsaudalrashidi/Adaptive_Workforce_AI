"""
Sandbox Diff Summary — Phase 43.

Computes what actually changed between the input file and the sandbox's
output file by reading BOTH files directly and comparing real data — never
by trusting whatever the LLM-generated code's own print output or the LLM's
own natural-language description claims it did. This is the same discipline
this session has applied everywhere else (the numeric guard in
ai_business_analyst_engine.py, the before/after test-count verification in
every phase report): a claim about what changed is only as good as an
independent check against the real artifact.

Supports the file types the sandbox's pre-bound libraries can produce:
  .xlsx  -> openpyxl, cell-level diff per sheet
  .csv   -> pandas, shape + cell-level diff
  .docx  -> python-docx, paragraph-level diff (difflib)
  .pptx  -> python-pptx, per-slide text-frame diff (difflib)

Unsupported extensions get an honest "diff not computed for this file
type" result rather than a fabricated summary.
"""
from __future__ import annotations

import difflib
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional


@dataclass
class CellChange:
    sheet: str
    cell: str
    old_value: Any
    new_value: Any


@dataclass
class DiffSummary:
    file_type: str
    supported: bool
    changed: bool
    summary_en: str
    summary_ar: str
    cell_changes_sample: List[Dict[str, Any]] = field(default_factory=list)
    n_cells_changed: int = 0
    n_rows_added: int = 0
    n_rows_removed: int = 0
    n_paragraphs_changed: int = 0
    n_paragraphs_added: int = 0
    n_paragraphs_removed: int = 0
    details: Dict[str, Any] = field(default_factory=dict)


_MAX_SAMPLE = 25


def compute_diff_summary(input_path: Path, output_path: Path) -> DiffSummary:
    suffix = output_path.suffix.lower()
    if suffix == ".xlsx":
        return _diff_xlsx(input_path, output_path)
    if suffix == ".csv":
        return _diff_csv(input_path, output_path)
    if suffix == ".docx":
        return _diff_docx(input_path, output_path)
    if suffix == ".pptx":
        return _diff_pptx(input_path, output_path)
    return DiffSummary(
        file_type=suffix or "unknown",
        supported=False,
        changed=True,  # unknown — cannot claim "no change" either
        summary_en=f"Diff not computed for file type '{suffix}' — the edited "
                   f"file is returned, but no independent before/after "
                   f"comparison is available for this format.",
        summary_ar=f"لم يُحتسب فرق للملف من نوع '{suffix}' — يُعاد الملف المُعدَّل، "
                   f"لكن لا تتوفر مقارنة مستقلة قبل/بعد لهذا التنسيق.",
    )


def _diff_xlsx(input_path: Path, output_path: Path) -> DiffSummary:
    import openpyxl

    wb_before = openpyxl.load_workbook(input_path, data_only=True)
    wb_after = openpyxl.load_workbook(output_path, data_only=True)

    cell_changes: List[CellChange] = []
    rows_added = 0
    rows_removed = 0

    all_sheets = list(dict.fromkeys(list(wb_before.sheetnames) + list(wb_after.sheetnames)))
    for sheet_name in all_sheets:
        ws_before = wb_before[sheet_name] if sheet_name in wb_before.sheetnames else None
        ws_after = wb_after[sheet_name] if sheet_name in wb_after.sheetnames else None

        max_row = max(ws_before.max_row if ws_before else 0, ws_after.max_row if ws_after else 0)
        max_col = max(ws_before.max_column if ws_before else 0, ws_after.max_column if ws_after else 0)

        if ws_before is not None and ws_after is not None:
            if ws_after.max_row > ws_before.max_row:
                rows_added += ws_after.max_row - ws_before.max_row
            elif ws_before.max_row > ws_after.max_row:
                rows_removed += ws_before.max_row - ws_after.max_row

        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                v_before = ws_before.cell(row=r, column=c).value if ws_before and r <= ws_before.max_row and c <= ws_before.max_column else None
                v_after = ws_after.cell(row=r, column=c).value if ws_after and r <= ws_after.max_row and c <= ws_after.max_column else None
                if v_before != v_after:
                    cell_ref = f"{sheet_name}!{ws_after.cell(row=r, column=c).coordinate if ws_after else ws_before.cell(row=r, column=c).coordinate}"
                    cell_changes.append(CellChange(sheet=sheet_name, cell=cell_ref, old_value=v_before, new_value=v_after))

    n_changed = len(cell_changes)
    changed = n_changed > 0 or rows_added > 0 or rows_removed > 0

    parts_en = []
    parts_ar = []
    if n_changed:
        parts_en.append(f"{n_changed} cell value(s) changed")
        parts_ar.append(f"تغيّرت قيمة {n_changed} خلية")
    if rows_added:
        parts_en.append(f"{rows_added} row(s) added")
        parts_ar.append(f"أُضيف {rows_added} صف")
    if rows_removed:
        parts_en.append(f"{rows_removed} row(s) removed")
        parts_ar.append(f"أُزيل {rows_removed} صف")
    summary_en = "; ".join(parts_en) if parts_en else "No differences detected between the original and edited file."
    summary_ar = "؛ ".join(parts_ar) if parts_ar else "لم يتم رصد أي فروقات بين الملف الأصلي والمُعدَّل."

    return DiffSummary(
        file_type=".xlsx",
        supported=True,
        changed=changed,
        summary_en=summary_en,
        summary_ar=summary_ar,
        cell_changes_sample=[
            {"sheet": cc.sheet, "cell": cc.cell, "old_value": _jsonable(cc.old_value), "new_value": _jsonable(cc.new_value)}
            for cc in cell_changes[:_MAX_SAMPLE]
        ],
        n_cells_changed=n_changed,
        n_rows_added=rows_added,
        n_rows_removed=rows_removed,
        details={"sheets_before": wb_before.sheetnames, "sheets_after": wb_after.sheetnames},
    )


def _diff_csv(input_path: Path, output_path: Path) -> DiffSummary:
    import pandas as pd

    df_before = pd.read_csv(input_path)
    df_after = pd.read_csv(output_path)

    cell_changes: List[Dict[str, Any]] = []
    n_changed = 0
    rows_added = max(0, len(df_after) - len(df_before))
    rows_removed = max(0, len(df_before) - len(df_after))

    common_cols = [c for c in df_before.columns if c in df_after.columns]
    common_rows = min(len(df_before), len(df_after))
    for i in range(common_rows):
        for col in common_cols:
            v_before = df_before.iloc[i][col]
            v_after = df_after.iloc[i][col]
            if pd.isna(v_before) and pd.isna(v_after):
                continue
            if v_before != v_after:
                n_changed += 1
                if len(cell_changes) < _MAX_SAMPLE:
                    cell_changes.append({"row": i, "column": col, "old_value": _jsonable(v_before), "new_value": _jsonable(v_after)})

    cols_added = [c for c in df_after.columns if c not in df_before.columns]
    cols_removed = [c for c in df_before.columns if c not in df_after.columns]

    changed = n_changed > 0 or rows_added > 0 or rows_removed > 0 or bool(cols_added) or bool(cols_removed)

    parts_en = []
    if n_changed:
        parts_en.append(f"{n_changed} cell value(s) changed")
    if rows_added:
        parts_en.append(f"{rows_added} row(s) added")
    if rows_removed:
        parts_en.append(f"{rows_removed} row(s) removed")
    if cols_added:
        parts_en.append(f"column(s) added: {', '.join(cols_added)}")
    if cols_removed:
        parts_en.append(f"column(s) removed: {', '.join(cols_removed)}")
    summary_en = "; ".join(parts_en) if parts_en else "No differences detected between the original and edited file."

    return DiffSummary(
        file_type=".csv", supported=True, changed=changed,
        summary_en=summary_en,
        summary_ar=summary_en,  # CSV column names are user data, not translated
        cell_changes_sample=cell_changes,
        n_cells_changed=n_changed, n_rows_added=rows_added, n_rows_removed=rows_removed,
        details={"columns_added": cols_added, "columns_removed": cols_removed},
    )


def _diff_docx(input_path: Path, output_path: Path) -> DiffSummary:
    import docx

    before_paras = [p.text for p in docx.Document(str(input_path)).paragraphs]
    after_paras = [p.text for p in docx.Document(str(output_path)).paragraphs]

    sm = difflib.SequenceMatcher(a=before_paras, b=after_paras)
    n_changed = n_added = n_removed = 0
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "replace":
            n_changed += max(i2 - i1, j2 - j1)
        elif tag == "insert":
            n_added += j2 - j1
        elif tag == "delete":
            n_removed += i2 - i1

    changed = n_changed > 0 or n_added > 0 or n_removed > 0
    parts_en = []
    if n_changed:
        parts_en.append(f"{n_changed} paragraph(s) changed")
    if n_added:
        parts_en.append(f"{n_added} paragraph(s) added")
    if n_removed:
        parts_en.append(f"{n_removed} paragraph(s) removed")
    summary_en = "; ".join(parts_en) if parts_en else "No differences detected between the original and edited document."

    return DiffSummary(
        file_type=".docx", supported=True, changed=changed,
        summary_en=summary_en, summary_ar=summary_en,
        n_paragraphs_changed=n_changed, n_paragraphs_added=n_added, n_paragraphs_removed=n_removed,
        details={"n_paragraphs_before": len(before_paras), "n_paragraphs_after": len(after_paras)},
    )


def _diff_pptx(input_path: Path, output_path: Path) -> DiffSummary:
    from pptx import Presentation

    def _slide_texts(path: Path) -> List[str]:
        prs = Presentation(str(path))
        texts = []
        for slide in prs.slides:
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
            texts.append("\n".join(parts))
        return texts

    before_slides = _slide_texts(input_path)
    after_slides = _slide_texts(output_path)

    sm = difflib.SequenceMatcher(a=before_slides, b=after_slides)
    n_changed = n_added = n_removed = 0
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "replace":
            n_changed += max(i2 - i1, j2 - j1)
        elif tag == "insert":
            n_added += j2 - j1
        elif tag == "delete":
            n_removed += i2 - i1

    changed = n_changed > 0 or n_added > 0 or n_removed > 0
    parts_en = []
    if n_changed:
        parts_en.append(f"{n_changed} slide(s) changed")
    if n_added:
        parts_en.append(f"{n_added} slide(s) added")
    if n_removed:
        parts_en.append(f"{n_removed} slide(s) removed")
    summary_en = "; ".join(parts_en) if parts_en else "No differences detected between the original and edited presentation."

    return DiffSummary(
        file_type=".pptx", supported=True, changed=changed,
        summary_en=summary_en, summary_ar=summary_en,
        details={"n_slides_before": len(before_slides), "n_slides_after": len(after_slides)},
    )


def _jsonable(value: Any) -> Any:
    """Best-effort conversion of a cell/dataframe value to something
    json.dumps can serialize (openpyxl/pandas can return datetime, Decimal,
    numpy scalar types, etc.)."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    try:
        import numpy as np
        if isinstance(value, np.generic):
            return value.item()
    except ImportError:
        pass
    return str(value)
