"""
Sandbox File Structure Extractor — Phase 43.

Reads the REAL structure of an uploaded file (sheet/column names, dtypes,
paragraph/slide counts) so the code-generation prompt is grounded in what
the file actually contains, rather than the LLM guessing plausible-sounding
column names. This mirrors the platform's established pattern (Phase 30+'s
"available"/"columns_detected" per-section flags across every BI module) of
never asking an engine — human-authored or LLM-generated — to act on data
it hasn't actually inspected.
"""
from __future__ import annotations

import io
from pathlib import Path
from typing import Union


def describe_file_structure(file_bytes: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".xlsx":
            return _describe_xlsx(file_bytes)
        if suffix == ".csv":
            return _describe_csv(file_bytes)
        if suffix == ".docx":
            return _describe_docx(file_bytes)
        if suffix == ".pptx":
            return _describe_pptx(file_bytes)
    except Exception as e:
        return f"(Could not inspect file structure: {e}. Filename: {filename})"
    return f"(Unsupported file type '{suffix}' for structure inspection. Filename: {filename})"


def _describe_xlsx(file_bytes: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    lines = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1), [])]
        lines.append(
            f"Sheet '{sheet_name}': {ws.max_row} rows x {ws.max_column} columns. "
            f"Header row: {header_row}"
        )
    return "Excel workbook (.xlsx):\n" + "\n".join(lines)


def _describe_csv(file_bytes: bytes) -> str:
    import pandas as pd
    df = pd.read_csv(io.BytesIO(file_bytes), nrows=5)
    full_len = sum(1 for _ in io.BytesIO(file_bytes).readlines()) - 1
    return (
        f"CSV file: {full_len} rows x {len(df.columns)} columns.\n"
        f"Columns and dtypes: {dict(df.dtypes.astype(str))}\n"
        f"First rows sample: {df.head(3).to_dict(orient='records')}"
    )


def _describe_docx(file_bytes: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(file_bytes))
    n_paragraphs = len(d.paragraphs)
    n_tables = len(d.tables)
    sample = [p.text for p in d.paragraphs[:5] if p.text.strip()]
    return (
        f"Word document (.docx): {n_paragraphs} paragraphs, {n_tables} tables.\n"
        f"First non-empty paragraphs sample: {sample}"
    )


def _describe_pptx(file_bytes: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(file_bytes))
    lines = [f"Presentation (.pptx): {len(prs.slides)} slides."]
    for i, slide in enumerate(prs.slides[:5]):
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        lines.append(f"Slide {i + 1}: {texts}")
    return "\n".join(lines)
