"""
Excel / CSV / Parquet Ingestion Engine — Phase 9.

Handles safe parsing of HR datasets in multiple formats with Arabic/English
header support, multi-sheet selection, formula stripping, and encoding detection.

Design rules:
  - Never execute cell formulas
  - Preserve original headers in metadata
  - Log every transformation
  - Detect language from column names
  - Handle Arabic column names and boolean values
  - Choose the most HR-relevant sheet when multiple sheets exist
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import numpy as np
import pandas as pd

logger = logging.getLogger(__name__)

# ── HR-relevance signal words ─────────────────────────────────────────────────

_HR_KEYWORDS_EN = {
    "employee", "emp", "staff", "worker", "person", "department", "dept",
    "salary", "attrition", "tenure", "hire", "job", "position", "performance",
    "manager", "training", "promotion", "leave", "absent", "resignation",
}

_HR_KEYWORDS_AR = {
    "موظف", "قسم", "راتب", "أداء", "مدير", "وظيفة", "تدريب", "استقالة",
    "الاستنزاف", "الدوران", "الراتب", "الخدمة",
}

_ARABIC_RE = re.compile(r"[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿]")

_BOOL_TRUE_AR  = {"نعم", "صح", "صحيح", "حقيقي"}
_BOOL_FALSE_AR = {"لا", "خطأ", "كاذب"}


# ── Return type ───────────────────────────────────────────────────────────────

@dataclass
class IngestionResult:
    success: bool
    df: Optional[pd.DataFrame]
    row_count: int
    column_count: int
    original_headers: list[str]
    normalized_headers: list[str]
    detected_language: str          # 'en' | 'ar' | 'mixed'
    detected_erp_style: Optional[str]
    encoding_used: str
    sheet_used: Optional[str]
    warnings: list[str] = field(default_factory=list)
    transformations: list[str] = field(default_factory=list)
    error: Optional[str] = None


# ── Engine ────────────────────────────────────────────────────────────────────

class ExcelIngestionEngine:
    """
    Safe ingestion of HR datasets in xlsx / xls / csv / parquet / json formats.

    Call:
        engine = ExcelIngestionEngine()
        result = engine.ingest(raw_bytes, filename)
    """

    SUPPORTED_EXTENSIONS = {".xlsx", ".xls", ".csv", ".parquet", ".json"}

    def ingest(self, raw_bytes: bytes, filename: str) -> IngestionResult:
        suffix = Path(filename).suffix.lower()
        if suffix not in self.SUPPORTED_EXTENSIONS:
            return IngestionResult(
                success=False, df=None, row_count=0, column_count=0,
                original_headers=[], normalized_headers=[],
                detected_language="unknown", detected_erp_style=None,
                encoding_used="none", sheet_used=None,
                error=f"Unsupported file type: {suffix!r}. Accepted: {sorted(self.SUPPORTED_EXTENSIONS)}",
            )

        if not raw_bytes:
            return IngestionResult(
                success=False, df=None, row_count=0, column_count=0,
                original_headers=[], normalized_headers=[],
                detected_language="unknown", detected_erp_style=None,
                encoding_used="none", sheet_used=None,
                error="File is empty.",
            )

        try:
            if suffix in (".xlsx", ".xls"):
                return self._load_excel(raw_bytes, filename)
            elif suffix == ".csv":
                return self._load_csv(raw_bytes, filename)
            elif suffix == ".parquet":
                return self._load_parquet(raw_bytes)
            elif suffix == ".json":
                return self._load_json(raw_bytes)
        except Exception as exc:
            logger.exception("Ingestion failed for %s", filename)
            return IngestionResult(
                success=False, df=None, row_count=0, column_count=0,
                original_headers=[], normalized_headers=[],
                detected_language="unknown", detected_erp_style=None,
                encoding_used="none", sheet_used=None,
                error=str(exc),
            )

        return IngestionResult(  # unreachable guard
            success=False, df=None, row_count=0, column_count=0,
            original_headers=[], normalized_headers=[],
            detected_language="unknown", detected_erp_style=None,
            encoding_used="none", sheet_used=None,
            error="Internal routing error.",
        )

    # ── Excel ─────────────────────────────────────────────────────────────────

    def _load_excel(self, raw_bytes: bytes, filename: str) -> IngestionResult:
        import openpyxl
        warnings: list[str] = []
        transformations: list[str] = []

        try:
            wb = openpyxl.load_workbook(
                io.BytesIO(raw_bytes),
                data_only=True,     # get computed values, not formulas
                read_only=False,
            )
        except Exception as exc:
            try:
                # xls fallback via pandas xlrd
                df = pd.read_excel(io.BytesIO(raw_bytes), engine="xlrd")
                return self._finalize_from_df(df, sheet_used="sheet1", warnings=warnings, transformations=transformations)
            except Exception:
                return IngestionResult(
                    success=False, df=None, row_count=0, column_count=0,
                    original_headers=[], normalized_headers=[],
                    detected_language="unknown", detected_erp_style=None,
                    encoding_used="utf-8", sheet_used=None,
                    error=f"Cannot open Excel file: {exc}",
                )

        sheet_names = wb.sheetnames
        if len(sheet_names) == 0:
            return IngestionResult(
                success=False, df=None, row_count=0, column_count=0,
                original_headers=[], normalized_headers=[],
                detected_language="unknown", detected_erp_style=None,
                encoding_used="utf-8", sheet_used=None,
                error="Excel file contains no sheets.",
            )

        sheet_name = self._choose_best_sheet(wb, sheet_names)
        if len(sheet_names) > 1:
            warnings.append(
                f"Multiple sheets detected ({sheet_names}). "
                f"Auto-selected '{sheet_name}' as most HR-relevant."
            )
            transformations.append(f"Sheet selected: {sheet_name!r}")

        ws = wb[sheet_name]

        # Detect hidden rows
        hidden_rows = sum(
            1 for row in ws.row_dimensions.values()
            if getattr(row, "hidden", False)
        )
        if hidden_rows > 0:
            warnings.append(f"{hidden_rows} hidden row(s) detected in sheet. These will be included.")

        # Build DataFrame from worksheet (data_only=True strips formulas)
        data = list(ws.values)
        if not data:
            return IngestionResult(
                success=False, df=None, row_count=0, column_count=0,
                original_headers=[], normalized_headers=[],
                detected_language="unknown", detected_erp_style=None,
                encoding_used="utf-8", sheet_used=sheet_name,
                error="Selected sheet is empty.",
            )

        # Find first non-empty row as header
        header_row_idx = 0
        for i, row in enumerate(data):
            if any(cell is not None and str(cell).strip() for cell in row):
                header_row_idx = i
                break

        if header_row_idx > 0:
            warnings.append(f"Skipped {header_row_idx} empty top row(s) before header.")
            transformations.append(f"Skipped {header_row_idx} leading empty rows")

        raw_headers = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(data[header_row_idx])]
        rows = data[header_row_idx + 1:]

        if not rows:
            return IngestionResult(
                success=False, df=None, row_count=0, column_count=0,
                original_headers=raw_headers, normalized_headers=raw_headers,
                detected_language=self._detect_language(raw_headers),
                detected_erp_style=None,
                encoding_used="utf-8", sheet_used=sheet_name,
                error="Sheet has headers but no data rows.",
            )

        df = pd.DataFrame(rows, columns=raw_headers)

        # Normalize Arabic booleans
        for col in df.select_dtypes(include="object").columns:
            series_lower = df[col].astype(str).str.strip().str.lower()
            if series_lower.isin(_BOOL_TRUE_AR | _BOOL_FALSE_AR | {"نعم", "لا"}).any():
                df[col] = df[col].map(
                    lambda v: True if str(v).strip() in _BOOL_TRUE_AR else
                              (False if str(v).strip() in _BOOL_FALSE_AR else v)
                )
                transformations.append(f"Normalized Arabic booleans in column '{col}'")

        return self._finalize_from_df(
            df, sheet_used=sheet_name, warnings=warnings, transformations=transformations
        )

    # ── CSV ───────────────────────────────────────────────────────────────────

    def _load_csv(self, raw_bytes: bytes, filename: str) -> IngestionResult:
        warnings: list[str] = []
        transformations: list[str] = []

        # Try multiple encodings in order
        encodings = ["utf-8-sig", "utf-8", "cp1256", "windows-1256", "latin-1"]
        df = None
        encoding_used = "unknown"

        for enc in encodings:
            try:
                df = pd.read_csv(io.BytesIO(raw_bytes), encoding=enc)
                encoding_used = enc
                if enc != "utf-8":
                    transformations.append(f"CSV decoded with encoding: {enc}")
                break
            except (UnicodeDecodeError, pd.errors.ParserError):
                continue

        if df is None:
            return IngestionResult(
                success=False, df=None, row_count=0, column_count=0,
                original_headers=[], normalized_headers=[],
                detected_language="unknown", detected_erp_style=None,
                encoding_used="unknown", sheet_used=None,
                error="CSV encoding not detected. Try saving as UTF-8.",
            )

        return self._finalize_from_df(
            df, sheet_used=None, warnings=warnings,
            transformations=transformations, encoding_used=encoding_used
        )

    # ── Parquet ───────────────────────────────────────────────────────────────

    def _load_parquet(self, raw_bytes: bytes) -> IngestionResult:
        try:
            df = pd.read_parquet(io.BytesIO(raw_bytes))
        except Exception as exc:
            return IngestionResult(
                success=False, df=None, row_count=0, column_count=0,
                original_headers=[], normalized_headers=[],
                detected_language="unknown", detected_erp_style=None,
                encoding_used="binary", sheet_used=None,
                error=f"Cannot parse parquet: {exc}",
            )
        return self._finalize_from_df(df, sheet_used=None, warnings=[], transformations=["Loaded from Parquet"], encoding_used="binary")

    # ── JSON ──────────────────────────────────────────────────────────────────

    def _load_json(self, raw_bytes: bytes) -> IngestionResult:
        try:
            df = pd.read_json(io.BytesIO(raw_bytes))
        except Exception as exc:
            return IngestionResult(
                success=False, df=None, row_count=0, column_count=0,
                original_headers=[], normalized_headers=[],
                detected_language="unknown", detected_erp_style=None,
                encoding_used="utf-8", sheet_used=None,
                error=f"Cannot parse JSON: {exc}",
            )
        return self._finalize_from_df(df, sheet_used=None, warnings=[], transformations=["Loaded from JSON"], encoding_used="utf-8")

    # ── Common finalization ───────────────────────────────────────────────────

    def _finalize_from_df(
        self,
        df: pd.DataFrame,
        sheet_used: Optional[str],
        warnings: list[str],
        transformations: list[str],
        encoding_used: str = "utf-8",
    ) -> IngestionResult:
        original_headers = list(df.columns.astype(str))

        # Handle duplicate columns
        if len(original_headers) != len(set(original_headers)):
            seen: dict[str, int] = {}
            new_cols = []
            for col in original_headers:
                if col in seen:
                    seen[col] += 1
                    new_cols.append(f"{col}_{seen[col]}")
                    warnings.append(f"Duplicate column '{col}' renamed to '{col}_{seen[col]}'")
                    transformations.append(f"Renamed duplicate column: {col} → {col}_{seen[col]}")
                else:
                    seen[col] = 0
                    new_cols.append(col)
            df.columns = new_cols

        # Normalize header strings (strip whitespace)
        normalized = [h.strip() for h in list(df.columns.astype(str))]
        if normalized != list(df.columns.astype(str)):
            transformations.append("Stripped whitespace from column names")
            df.columns = normalized

        # Parse salary columns with currency/comma formatting
        for col in df.columns:
            col_lower = col.lower()
            if any(kw in col_lower for kw in ("salary", "راتب", "income", "pay", "compensation")):
                if df[col].dtype == object:
                    cleaned = df[col].astype(str).str.replace(r"[, \s]", "", regex=True)
                    cleaned = cleaned.str.replace(r"[^\d.\-]", "", regex=True)
                    try:
                        df[col] = pd.to_numeric(cleaned, errors="coerce")
                        transformations.append(f"Normalized salary column '{col}' (removed currency/commas)")
                    except Exception:
                        pass

        detected_lang = self._detect_language(list(df.columns.astype(str)))
        detected_erp = self._detect_erp_style(list(df.columns.astype(str)))

        return IngestionResult(
            success=True,
            df=df,
            row_count=len(df),
            column_count=len(df.columns),
            original_headers=original_headers,
            normalized_headers=list(df.columns.astype(str)),
            detected_language=detected_lang,
            detected_erp_style=detected_erp,
            encoding_used=encoding_used,
            sheet_used=sheet_used,
            warnings=warnings,
            transformations=transformations,
        )

    # ── Helper: sheet selection ───────────────────────────────────────────────

    def _choose_best_sheet(self, wb: "openpyxl.Workbook", sheet_names: list[str]) -> str:
        if len(sheet_names) == 1:
            return sheet_names[0]

        best_sheet = sheet_names[0]
        best_score = -1

        for name in sheet_names:
            score = 0
            name_lower = name.lower()
            # Prefer sheets named like HR data
            hr_sheet_keywords = {"employee", "emp", "hr", "staff", "data", "موظف", "البيانات"}
            for kw in hr_sheet_keywords:
                if kw in name_lower:
                    score += 3

            # Check how many HR-like column names the sheet has
            ws = wb[name]
            headers = []
            for cell in next(ws.iter_rows(max_row=1, values_only=True), []):
                if cell is not None:
                    headers.append(str(cell).lower())

            for h in headers:
                if any(kw in h for kw in _HR_KEYWORDS_EN):
                    score += 1
                if any(kw in h for kw in _HR_KEYWORDS_AR):
                    score += 1

            if score > best_score:
                best_score = score
                best_sheet = name

        return best_sheet

    # ── Helper: language detection ────────────────────────────────────────────

    def _detect_language(self, headers: list[str]) -> str:
        arabic_count = sum(1 for h in headers if _ARABIC_RE.search(h))
        english_count = sum(1 for h in headers if re.search(r"[a-zA-Z]", h))

        if arabic_count == 0 and english_count == 0:
            return "unknown"
        if arabic_count > 0 and english_count > 0:
            return "mixed"
        if arabic_count > english_count:
            return "ar"
        return "en"

    # ── Helper: ERP detection ─────────────────────────────────────────────────

    def _detect_erp_style(self, headers: list[str]) -> Optional[str]:
        headers_lower = " ".join(h.lower() for h in headers)

        if any(kw in headers_lower for kw in ("workday", "wd_", "worker_id")):
            return "Workday"
        if any(kw in headers_lower for kw in ("successfactors", "sf_", "ecf_")):
            return "SAP SuccessFactors"
        if any(kw in headers_lower for kw in ("oracle", "hcm_", "person_number")):
            return "Oracle HCM"
        if any(kw in headers_lower for kw in ("mudad", "مدد", "الراتب_الأساسي")):
            return "Mudad (Saudi Payroll)"
        if any(kw in headers_lower for kw in ("qiwa", "qiwa_", "saudization")):
            return "Qiwa (Saudi Labour)"
        if any(kw in headers_lower for kw in ("bamboohr", "bamboo")):
            return "BambooHR"

        return None
